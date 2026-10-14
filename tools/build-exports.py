#!/usr/bin/env python3
"""Build the downloadable files and the static web page from data/cmt-genes.csv.

Writes to dist/ (not committed):

  dist/cmt-genes.csv        the table, unchanged
  dist/cmt-genes.xlsx       the table, a column guide and an About sheet
  dist/cmt-genes.json       one object per row
  dist/datapackage.json     Frictionless Data Package describing every column
  dist/index.html           the GitHub Pages site: searchable table and downloads

The column descriptions come from the column table in README.md, so the
README stays the single place to edit them.

Usage: python3 tools/build-exports.py [--version v2026.10]
The version defaults to the VERSION environment variable, then to "dev".
"""
import csv, datetime, html, json, os, re, shutil, sys

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DIST = os.path.join(ROOT, "dist")
REPO = "https://github.com/dknauss/cmt-gene-table"
PAGES = "https://dknauss.github.io/cmt-gene-table/"
BLOG = "https://cmt.blog/cmt-gene-table/"
TITLE = "Open CMT Gene Table"
CREDIT = "Open CMT Gene Table (cmt.blog), using ClinGen (CC0) and Mondo (CC BY 4.0)"
INH = {"AD": "Dominant", "AR": "Recessive", "SD": "Semidominant", "XL": "X-linked"}


def version():
    if "--version" in sys.argv:
        return sys.argv[sys.argv.index("--version") + 1]
    return os.environ.get("VERSION") or "dev"


def column_guide(fields):
    """Read the README column table: {column: (source, description)}."""
    guide = {}
    text = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    for line in text.splitlines():
        m = re.match(r"^\|\s*(`[^|]+`)\s*\|\s*([^|]+?)\s*\|\s*(.+?)\s*\|$", line)
        if not m:
            continue
        names = re.findall(r"`([^`]+)`", m.group(1))
        for n in names:
            desc = m.group(3).replace("`; `", "semicolons").replace(" below", " in the README")
            guide[n] = (m.group(2).strip(), re.sub(r"`([^`]+)`", r"\1", desc).strip())
    missing = [f for f in fields if f not in guide]
    if missing:
        sys.exit("README column table is missing: " + ", ".join(missing))
    return guide


def split_links(value):
    """'Label: https://…; Label: https://…' -> [(label, url)]."""
    out = []
    for part in [p.strip() for p in value.split("; ") if p.strip()]:
        m = re.match(r"^(.*?):\s+(https?://\S+)$", part)
        out.append((m.group(1), m.group(2)) if m else (part, None))
    return out


def links_html(value):
    items = []
    for label, url in split_links(value):
        label = html.escape(label)
        items.append(f'<a href="{html.escape(url)}">{label}</a>' if url else label)
    return "; ".join(items)


def names(value):
    return [x.strip() for x in value.split(";") if x.strip()]


def build_xlsx(rows, fields, guide, ver, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "CMT genes"
    ws.append(fields)
    for r in rows:
        ws.append([r[f] for f in fields])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, f in enumerate(fields, 1):
        longest = max([len(f)] + [len(r[f]) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=len(str(c.value or "")) > 60)

    cg = wb.create_sheet("Columns")
    cg.append(["column", "source", "description"])
    for f in fields:
        cg.append([f, guide[f][0], guide[f][1]])
    for c in cg[1]:
        c.font = Font(bold=True)
    cg.column_dimensions["A"].width = 28
    cg.column_dimensions["B"].width = 26
    cg.column_dimensions["C"].width = 100
    for row in cg.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    ab = wb.create_sheet("About")
    for line in [
        [TITLE],
        ["Version", ver],
        ["Built", datetime.date.today().isoformat()],
        ["Rows", f"{len(rows)} ClinGen curations covering {len({r['gene'] for r in rows})} genes"],
        ["Explained and footnoted at", BLOG],
        ["Source code, history and every hand-checked source", REPO],
        ["Licence", "Data: CC BY 4.0. Credit: " + CREDIT],
        ["Not medical advice", "A registered trial is not a promising or available treatment. Talk to a clinician or genetic counsellor about any result."],
    ]:
        ab.append(line)
    ab["A1"].font = Font(bold=True, size=14)
    ab.column_dimensions["A"].width = 44
    ab.column_dimensions["B"].width = 110
    wb.save(path)


def build_datapackage(rows, fields, guide, ver):
    types = {"classification_date": "date"}
    return {
        "name": "open-cmt-gene-table",
        "title": TITLE,
        "description": "Genes linked to Charcot-Marie-Tooth disease and related inherited neuropathies, one row per ClinGen gene-disease curation, with subtype names, sources and plain-language notes.",
        "version": ver,
        "created": datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "homepage": PAGES,
        "licenses": [{"name": "CC-BY-4.0", "title": "Creative Commons Attribution 4.0", "path": "https://creativecommons.org/licenses/by/4.0/"}],
        "contributors": [{"title": "Dan Knauss", "role": "author", "path": "https://cmt.blog/"}],
        "sources": [
            {"title": "ClinGen Gene-Disease Validity", "path": "https://search.clinicalgenome.org/kb/gene-validity"},
            {"title": "Mondo Disease Ontology via the Monarch Initiative", "path": "https://monarchinitiative.org/"},
            {"title": "GeneReviews", "path": "https://www.ncbi.nlm.nih.gov/books/NBK1116/"},
            {"title": "Genomics England PanelApp and PanelApp Australia", "path": "https://panelapp.genomicsengland.co.uk/"},
            {"title": "ClinicalTrials.gov", "path": "https://clinicaltrials.gov/"},
        ],
        "resources": [{
            "name": "cmt-genes",
            "path": "cmt-genes.csv",
            "format": "csv",
            "mediatype": "text/csv",
            "encoding": "utf-8",
            "schema": {
                "fields": [{"name": f, "type": types.get(f, "string"), "description": f"{guide[f][1]} (Source: {guide[f][0]}.)"} for f in fields],
                "primaryKey": ["gene", "mondo_id"],
            },
        }],
    }


def build_page(rows, fields, guide, ver, sizes):
    body = []
    for r in rows:
        main = names(r["curation_subtype_names"])
        other = names(r["other_gene_names"])
        older = names(r["unverified_names"])
        cell = html.escape(", ".join(main)) if main else "—"
        extra = []
        if other:
            extra.append("Other names, not rated here: " + html.escape(", ".join(other)) + ".")
        if older:
            extra.append("Older or unverified: " + html.escape(", ".join(older)) + ".")
        if extra:
            cell += '<br><span class="minor">' + " ".join(extra) + "</span>"
        more = [f"<p>{html.escape(r['plain_note'])}</p>"] if r["plain_note"] else []
        for label, key in [("Trials", "therapy_programs"), ("Trials in other conditions", "other_condition_trials"), ("GeneReviews", "genereviews"), ("PanelApp", "panelapp")]:
            if r[key]:
                more.append(f"<p><strong>{label}:</strong> {links_html(r[key])}</p>")
        details = f"<details><summary>More<span class=\"sr\"> about {html.escape(r['gene'])}</span></summary>{''.join(more)}</details>" if more else ""
        search = " ".join([r["gene"], r["hgnc_id"], r["mondo_id"], r["clingen_disease_label"], r["historical_subtype_names"], r["unverified_names"], r["phenotype"], INH.get(r["inheritance"], r["inheritance"])])
        body.append(
            f'<tr data-search="{html.escape(search.lower())}">'
            f'<th scope="row">{html.escape(r["gene"])}</th>'
            f"<td>{cell}</td>"
            f"<td>{html.escape(r['clingen_disease_label'])}<br><span class=\"minor\">{html.escape(r['mondo_id'])}</span></td>"
            f"<td>{INH.get(r['inheritance'], html.escape(r['inheritance']))}</td>"
            f'<td data-sort="{html.escape(r["clingen_classification"])}"><a href="{html.escape(r["clingen_report"])}">{html.escape(r["clingen_classification"])}</a> ({r["classification_date"][:4]})</td>'
            f"<td>{html.escape(r['phenotype'].replace(';', ','))}</td>"
            f"<td>{details}</td>"
            "</tr>"
        )
    aliases = "{}"
    ap = os.path.join(ROOT, "site", "gene-aliases.json")
    if os.path.exists(ap):
        aliases = json.dumps(json.load(open(ap, encoding="utf-8")), separators=(",", ":"))
    genes = len({r["gene"] for r in rows})
    kb = lambda n: f"{max(1, round(sizes[n] / 1024))} KB"
    tpl = open(os.path.join(ROOT, "site", "pages", "index.html"), encoding="utf-8").read()
    columns_html = "\n".join(f"<tr><th scope=\"row\"><code>{html.escape(f)}</code></th><td>{html.escape(guide[f][0])}</td><td>{html.escape(guide[f][1])}</td></tr>" for f in fields)
    return (tpl.replace("{{VERSION}}", html.escape(ver))
               .replace("{{BUILT}}", datetime.date.today().isoformat())
               .replace("{{ROWS}}", str(len(rows)))
               .replace("{{GENES}}", str(genes))
               .replace("{{TBODY}}", "\n".join(body))
               .replace("{{ALIASES}}", aliases)
               .replace("{{COLUMNS}}", columns_html)
               .replace("{{REPO}}", REPO)
               .replace("{{BLOG}}", BLOG)
               .replace("{{CREDIT}}", html.escape(CREDIT))
               .replace("{{SIZE_XLSX}}", kb("cmt-genes.xlsx"))
               .replace("{{SIZE_CSV}}", kb("cmt-genes.csv"))
               .replace("{{SIZE_JSON}}", kb("cmt-genes.json")))


def main():
    ver = version()
    src = os.path.join(ROOT, "data", "cmt-genes.csv")
    with open(src, encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        fields = reader.fieldnames
        rows = list(reader)
    guide = column_guide(fields)

    shutil.rmtree(DIST, ignore_errors=True)
    os.makedirs(DIST)
    shutil.copyfile(src, os.path.join(DIST, "cmt-genes.csv"))
    build_xlsx(rows, fields, guide, ver, os.path.join(DIST, "cmt-genes.xlsx"))
    with open(os.path.join(DIST, "cmt-genes.json"), "w", encoding="utf-8") as fh:
        json.dump(rows, fh, ensure_ascii=False, indent=1)
        fh.write("\n")
    with open(os.path.join(DIST, "datapackage.json"), "w", encoding="utf-8") as fh:
        json.dump(build_datapackage(rows, fields, guide, ver), fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    sizes = {n: os.path.getsize(os.path.join(DIST, n)) for n in ("cmt-genes.csv", "cmt-genes.xlsx", "cmt-genes.json")}
    with open(os.path.join(DIST, "index.html"), "w", encoding="utf-8") as fh:
        fh.write(build_page(rows, fields, guide, ver, sizes))
    open(os.path.join(DIST, ".nojekyll"), "w").close()
    print(f"{len(rows)} rows, version {ver}: " + ", ".join(sorted(os.listdir(DIST))))


if __name__ == "__main__":
    main()
